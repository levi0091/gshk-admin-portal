"""One-off: CR's TPSI worksheet -> committed nd2a_schema.json / nd2b_schema.json.

Source: tests/fixtures/cr-examples/Worksheet in TPSI API Interface v1.0.14.xlsx,
sheets "ND2A" and "ND2B" — CR's XML parameter contract, laid out exactly as the
NAR1 sheet `gen_nar1_schema.py` reads: the parameter name sits in one of the
columns A..H and the column index IS the nesting depth; then Type, Mandatory,
Length, Remark.

As for NAR1, where the worksheet and CR's own example instances disagree, THE
EXAMPLES WIN: they are what CR's server produced. Every fixup below is evidenced
by `tests/fixtures/cr-examples/validateForm/validate_ND2*.xml`, and the
cross-check at the bottom fails if any element an example uses is missing.

    cd backend && uv run python scripts/gen_nd2_schema.py
"""
import json
import pathlib
import re
import sys
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

HERE = pathlib.Path(__file__).resolve().parent.parent
SRC = HERE / "tests/fixtures/cr-examples/Worksheet in TPSI API Interface v1.0.14.xlsx"
OUT = HERE / "services/officer_changes/schema"
EXAMPLES = HERE / "tests/fixtures/cr-examples/validateForm"

_NAME_FIXUPS = {"Eform": "EForm"}

#: Elements the examples carry that the worksheet's bean listing omits:
#: (form, bean, insert_after, node). `associatedCapacityDesc` sits inside every
#: body-corporate DIRECTOR appointment in CR's examples (the consent signer's
#: capacity); the worksheet lists it only at form level.
_INSERT_AFTER = [
    ("ND2A", "appOfBcBean", "associatedPersonName",
     {"name": "associatedCapacityDesc", "data_type": "String", "mandatory": False,
      "max_length": 500, "remark": "Consent signer's capacity (from CR's examples)"}),
]


def _cell(row, i):
    return "" if i >= len(row) or row[i] is None else str(row[i]).strip()


def _columns(header):
    names = [_cell(header, i) for i in range(len(header))]
    type_col = names.index("Type")
    return type_col, type_col + 1, type_col + 2, type_col + 3


def build(sheet: str) -> dict:
    ws = load_workbook(SRC, read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    type_col, mand_col, len_col, remark_col = _columns(rows[0])
    root = {"name": "", "depth": -2, "children": []}
    stack = [root]
    for row in rows[1:]:
        depth = next((i for i in range(type_col) if _cell(row, i)), None)
        if depth is None:
            continue
        name = _NAME_FIXUPS.get(_cell(row, depth), _cell(row, depth))
        length = _cell(row, len_col)
        node = {
            "name": name, "depth": depth,
            "data_type": _cell(row, type_col),
            "mandatory": _cell(row, mand_col).upper() == "Y",
            "max_length": int(float(length)) if re.fullmatch(r"\d+(\.0)?", length) else None,
            "remark": _cell(row, remark_col), "children": [],
        }
        while stack[-1]["depth"] >= depth:
            stack.pop()
        stack[-1]["children"].append(node)
        stack.append(node)
    tree = root["children"][0]
    for form, bean, after, extra in _INSERT_AFTER:
        if form != sheet:
            continue
        holder = _find_named(tree, bean)
        index = next(i for i, c in enumerate(holder["children"]) if c["name"] == after)
        holder["children"].insert(index + 1, {**extra, "depth": holder["depth"] + 1,
                                              "children": []})
    return tree


def _find_named(node, name):
    if node["name"] == name:
        return node
    for child in node["children"]:
        hit = _find_named(child, name)
        if hit:
            return hit
    return None


def _names(node, out):
    out.add(node["name"])
    for child in node["children"]:
        _names(child, out)
    return out


def _example_names(path: pathlib.Path) -> set[str]:
    # CR's own "Cease Corporate Director" example has a stray "<" before a
    # closing tag; read it as CR evidently meant it.
    text = path.read_text(encoding="utf8").replace("<</", "</")
    names = set()
    for el in ET.fromstring(text).iter():
        names.add(el.tag.split("}")[-1])
    return names - {"Envelope", "Body", "validateForm"}


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    failures = []
    for sheet, code in (("ND2A", "nd2a"), ("ND2B", "nd2b")):
        tree = build(sheet)
        (OUT / f"{code}_schema.json").write_text(
            json.dumps(tree, ensure_ascii=False, indent=2) + "\n", encoding="utf8")
        known = _names(tree, set())
        for example in sorted(EXAMPLES.glob(f"validate_{sheet}(*.xml")):
            missing = _example_names(example) - known
            if missing:
                failures.append(f"{example.name}: {sorted(missing)}")
    for failure in failures:
        print("MISSING", failure, file=sys.stderr)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
